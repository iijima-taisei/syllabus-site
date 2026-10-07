#!/usr/bin/env python3
"""Moodle「令和８年度・基幹教育科目履修関係ページ」のクラス割(Excel)から、
クラス記号 → 担当の講義コード の対応だけを作り data/classwari.json に保存する。
使い方: python3 import_classwari.py <クラス割.xlsx> ...   (openpyxl が必要。通信なし)
対応するファイル(ファイル名で判別):
  - 学術英語・アカデミックリテラシー / アカデミックライティング のクラス割 (シート=曜限、「担当：教員」の下に学生とクラス)
  - ≪物理科目≫クラス割(基礎・概論) (行=学生、列=科目ごとの担当教員)
- 学籍番号・氏名などの個人の情報は読み捨て、出力には入れない。
- クラスの学生全員が同じ担当になる場合だけ出力する(個人単位で分かれるクラスは対象外)。
- 担当を B表(data/btable.json)の講義と突き合わせて講義コードにする。"""
import collections, json, pathlib, re, sys, unicodedata
import openpyxl

ROOT = pathlib.Path(__file__).parent
DAYS = "月火水木金"
BT = json.loads((ROOT / "data" / "btable.json").read_text(encoding="utf-8"))["rows"]


def nz(x):
    return re.sub(r"[\s,，.．・]", "", unicodedata.normalize("NFKC", str(x or ""))).lower()


def same_teacher(a, b):
    a, b = nz(a), nz(b)
    return bool(a and b) and (a[:3] == b[:3] or a in b or b in a)


def pick(cands, teacher, cls):
    hit = [k for k in cands if same_teacher(BT[k]["teacher"], teacher)]
    if len(hit) > 1:
        hit = [k for k in hit if cls in BT[k]["cs"]] or hit
    return hit[0] if len(hit) == 1 else None


def english(path, subject):
    """シート=曜限(学期)、担当ごとに学生が並ぶ形式"""
    agg = collections.defaultdict(collections.Counter)
    for ws in openpyxl.load_workbook(path, read_only=True, data_only=True).worksheets:
        t = nz(ws.title)
        m = re.search(rf"([{DAYS}])(\d)", t)
        sem = "秋学期" if "秋" in t else "冬学期" if "冬" in t else "後期"
        if not m:
            continue
        teacher = None
        for r in ws.iter_rows(values_only=True):
            c0 = nz(r[0]) if r else ""
            if c0.startswith("担当"):
                teacher = re.split(r"[:：]", unicodedata.normalize("NFKC", str(r[0])), maxsplit=1)[-1]
                continue
            cls = next((nz(c).upper() for c in r if re.fullmatch(r"[ils]1-\d+", nz(c))), None)
            if cls and teacher:
                agg[cls][(sem, m.group(1), int(m.group(2)), teacher)] += 1
    out = {}
    for cls, cnt in agg.items():
        if len(cnt) != 1:
            continue
        sem, d, p, teacher = next(iter(cnt))
        cands = [k for k, v in BT.items() if subject in v["name"] and v["sem"] == sem and d in v["day"] and p in v["per"]]
        code = pick(cands, teacher, cls)
        if code:
            out[cls] = code
        else:
            print(f"  {subject}: B表と対応づけられません {cls} {sem}{d}{p}")
    return out, len(agg)


def physics(path):
    """行=学生、列=科目(力学基礎 など)ごとの担当教員"""
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    head = [unicodedata.normalize("NFKC", str(c or "")).strip() for c in rows[0]]
    ci = head.index("クラス")
    agg = collections.defaultdict(lambda: collections.defaultdict(collections.Counter))
    for r in rows[1:]:
        cls = nz(r[ci]).upper()
        if not re.fullmatch(r"[ILS]1-\d+", cls):
            continue
        for i, subj in enumerate(head):
            if subj.endswith(("基礎", "概論")) and r[i]:
                agg[subj][cls][unicodedata.normalize("NFKC", str(r[i])).strip()] += 1
    out, total = collections.defaultdict(list), 0
    for subj, d in agg.items():
        for cls, cnt in d.items():
            total += 1
            if len(cnt) != 1:
                continue
            teacher = next(iter(cnt))
            cands = [k for k, v in BT.items() if v["name"].strip() == subj]
            code = pick(cands, teacher.split("・")[0], cls)
            if code:
                out[cls].append(code)
            else:
                print(f"  {subj}: B表と対応づけられません {cls} {teacher}")
    return out, total


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    res, notes = collections.defaultdict(lambda: {"req": set(), "sec": set()}), []
    for f in sys.argv[1:]:
        name = pathlib.Path(f).name
        if "アカデミックリテラシー" in name:
            m, n = english(f, "アカデミックリテラシー")
            items, kind = {k: [v] for k, v in m.items()}, "req"
        elif "アカデミックライティング" in name:
            m, n = english(f, "アカデミックライティング")
            items, kind = {k: [v] for k, v in m.items()}, "req"
        elif "物理" in name:
            items, n = physics(f)
            kind = "sec"
        else:
            print("対応していないファイル:", name)
            continue
        for k, v in items.items():
            res[k][kind].update(v)
        notes.append(f"{name}: {sum(len(v) for v in items.values())}件 / {n}")
        print(notes[-1])
    # req = 必修(学術英語) / sec = 履修するなら担当が決まっている授業(物理。クラスによっては選択科目)
    out = {k: {t: sorted(v[t]) for t in ("req", "sec") if v[t]} for k, v in sorted(res.items())}
    bad = [(k, c) for k, d in out.items() for cs in d.values() for c in cs if k not in BT[c]["cs"] and not BT[c]["fl"]]
    if bad:
        print("注意: B表の対象クラスに含まれない対応:", bad)
    (ROOT / "data" / "classwari.json").write_text(json.dumps({
        "source": "Moodle「令和８年度・基幹教育科目履修関係ページ」の後期クラス割(2026/9/28更新)のうち、クラス単位で担当が決まるもの(学術英語リテラシー・ライティング、物理の基礎・概論)",
        "classes": out}, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"data/classwari.json: {len(out)}クラス / 必修 {sum(len(v.get('req', [])) for v in out.values())} / 担当指定 {sum(len(v.get('sec', [])) for v in out.values())}")


if __name__ == "__main__":
    main()
